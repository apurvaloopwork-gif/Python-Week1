from openpyxl import workbook,load_workbook

wb=load_workbook('employees.xlsx')
ws=wb.active

# print(ws['B2'].value)

# #Change the value 

# ws['B2']="Lily"
# print(ws['B2'].value)


#to create new sheet 
wb.create_sheet("sheet1")
print(wb.sheetnames) #get all sheets 
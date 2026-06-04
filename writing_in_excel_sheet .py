from openpyxl import Workbook

wb = Workbook()
ws = wb.active      
ws.title = "Employees"

ws["A1"] = "ID"
ws["B1"] = "Name"
ws["C1"] = "Department"


ws.append([1, "Apurva", "IT"])
ws.append([2, "Peter", "Marketing"])
ws.append([3, "Tony", "HR"])
ws.append([4, "Shelby", "Accounts"])


wb.save("employees.xlsx")

print("Excel file created successfully.")